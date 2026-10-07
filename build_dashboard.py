"""Build DC_Site_Dashboard.html from DC_Site_Scoring_Model.xlsx (read-only).

Run:  .venv/bin/python build_dashboard.py
Reads cached values from the workbook, embeds them as JSON into
dashboard_template.html, and checks the scoring engine against the
workbook's own Sensitivity sheet before writing the output.
"""
import json
import math
from pathlib import Path

import openpyxl

ROOT = Path(__file__).parent
XLSX = ROOT / "DC_Site_Scoring_Model.xlsx"
TEMPLATE = ROOT / "dashboard_template.html"
OUT = ROOT / "DC_Site_Dashboard.html"

wb = openpyxl.load_workbook(XLSX, data_only=True)


def num(x):
    return float(x) if isinstance(x, (int, float)) else x


# ---- Weights ---------------------------------------------------------------
wsW = wb["Weights"]
scen_names = [wsW.cell(4, c).value for c in range(3, 9)]  # Base..Speed
categories = [wsW.cell(r, 2).value for r in range(5, 12)]
scenarios = {
    n: [num(wsW.cell(r, 3 + i).value) for r in range(5, 12)]
    for i, n in enumerate(scen_names)
}
share_in_cat = {wsW.cell(r, 1).value: num(wsW.cell(r, 4).value) for r in range(16, 34)}

# ---- Data ------------------------------------------------------------------
wsD = wb["Data"]
markets = [wsD.cell(4, c).value for c in range(8, 14)]
variables = []
for r in range(5, 23):
    v = {
        "id": wsD.cell(r, 1).value,
        "cat": wsD.cell(r, 2).value,
        "name": wsD.cell(r, 3).value,
        "unit": wsD.cell(r, 4).value,
        "type": wsD.cell(r, 5).value,
        "better": wsD.cell(r, 6).value,
        "transform": wsD.cell(r, 7).value,
        "values": [num(wsD.cell(r, c).value) for c in range(8, 14)],
        "confidence": wsD.cell(r, 14).value,
        "sources": wsD.cell(r, 15).value,
        "notes": wsD.cell(r, 16).value,
        "judgment": num(wsD.cell(r, 17).value) or 0,
        "share": share_in_cat[wsD.cell(r, 1).value],
    }
    variables.append(v)

# ---- Facts, latency, sources -----------------------------------------------
wsF = wb["MarketFacts"]
facts = [
    {"topic": wsF.cell(r, 1).value, "by": [wsF.cell(r, c).value for c in range(2, 8)]}
    for r in range(5, 12)
]
wsL = wb["Latency"]
latency = {
    "route_factor": num(wsL["B4"].value),
    "speed": num(wsL["B5"].value),
    "metros": [
        {"name": wsL.cell(r, 1).value, "pop": num(wsL.cell(r, 4).value),
         "rtt": [num(wsL.cell(r, c).value) for c in range(6, 12)]}
        for r in range(18, 31)
    ],
}
wsS = wb["Sources"]
sources = [
    {"id": wsS.cell(r, 1).value, "source": wsS.cell(r, 2).value, "url": wsS.cell(r, 3).value,
     "reliability": wsS.cell(r, 4).value, "used": wsS.cell(r, 5).value}
    for r in range(5, wsS.max_row + 1) if wsS.cell(r, 1).value
]

# ---- Reference results straight from the workbook (for the in-page check) --
wsSe = wb["Sensitivity"]
ref_scores = {
    scen_names[i]: [num(wsSe.cell(r, 2 + i).value) for r in range(5, 11)]
    for i in range(6)
}
# workbook Data/Summary order is the same market order used here
ref_sweep = [[num(wsSe.cell(r, c).value) for c in range(1, 8)] for r in range(24, 35)]
ref_grid = [[num(wsSe.cell(r, c).value) for c in range(1, 9)] for r in range(39, 55)]

data = {
    "markets": markets, "categories": categories, "scenarios": scenarios,
    "variables": variables, "facts": facts, "latency": latency, "sources": sources,
    "ref": {"scores": ref_scores, "sweep": ref_sweep, "grid": ref_grid},
}


# ---- Python mirror of the JS engine, used as a build-time assertion --------
def normalize():
    out = []
    for v in variables:
        vals = v["values"]
        if v["type"] == "rubric":
            s = [((x - 1) / 4 if v["better"] == "higher" else (5 - x) / 4) * 100 for x in vals]
        else:
            t = [math.log(x) if v["transform"] == "log" else x for x in vals]
            lo, hi = min(t), max(t)
            if hi == lo:
                s = [50] * 6
            else:
                s = [((x - lo) / (hi - lo) if v["better"] == "higher" else (hi - x) / (hi - lo)) * 100
                     for x in t]
        out.append(s)
    return out


NORM = normalize()
SOURCED_MAX_JUDGMENT = 0.34  # matches the memo's DFW 87 / Seattle 75 / next best 44
data["sourcedMax"] = SOURCED_MAX_JUDGMENT


def score(cat_w, sourced_only=False):
    tot = [0.0] * 6
    wsum = 0.0
    for v, s in zip(variables, NORM):
        w = cat_w[categories.index(v["cat"])] * v["share"]
        if sourced_only and v["judgment"] > SOURCED_MAX_JUDGMENT:
            continue
        wsum += w
        for i in range(6):
            tot[i] += w * s[i]
    return [t / wsum for t in tot]


if __name__ == "__main__":
    # 1) scenario scores match the workbook
    for n in scen_names:
        got = score(scenarios[n])
        for g, r in zip(got, ref_scores[n]):
            assert abs(g - r) < 0.05, (n, g, r)
    print("scenario scores match workbook")
    # 2) one-way sweep match (power weight p, other categories scaled pro rata)
    base = scenarios["Base"]
    for row in ref_sweep:
        p = row[0]
        k = (1 - p) / (1 - base[0])
        w = [p] + [x * k for x in base[1:]]
        for g, r in zip(score(w), row[1:]):
            assert abs(g - r) < 0.05, ("sweep", p, g, r)
    print("one-way sweep matches workbook")
    # 3) two-way grid match
    for row in ref_grid:
        p, l = row[0], row[1]
        k = (1 - p - l) / (1 - base[0] - base[2])
        w = [p, base[1] * k, l, base[3] * k, base[4] * k, base[5] * k, base[6] * k]
        for g, r in zip(score(w), row[2:8]):
            assert abs(g - r) < 0.05, ("grid", p, l, g, r)
    print("two-way grid matches workbook")
    # 4) judgment share and sourced-only view
    jb = sum(base[categories.index(v["cat"])] * v["share"] * v["judgment"] for v in variables)
    print("judgment share of base weight: %.1f%%" % (jb * 100))
    print("sourced-only:", dict(zip(markets, [round(x, 1) for x in score(base, True)])))

    html = TEMPLATE.read_text().replace("/*__DATA__*/null", json.dumps(data, default=str))
    OUT.write_text(html)
    print("wrote", OUT, len(html) // 1024, "KB")
